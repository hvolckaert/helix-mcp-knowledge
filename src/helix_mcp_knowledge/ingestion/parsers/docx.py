"""DOCX parser for headings, paragraphs and tables."""

from pathlib import Path

from docx import Document as DocxDocument

from ...models.chunk import ChunkType
from ..normalizer import normalize_text
from .base import ParsedBlock, ParsedDocument


class DocxParser:
    def parse(self, path: Path) -> ParsedDocument:
        document = DocxDocument(path)
        title = normalize_text(document.core_properties.title or "")
        if not title or title.casefold() in {"document", "microsoft word", "word document"}:
            title = path.stem
        blocks: list[ParsedBlock] = []
        for paragraph in document.paragraphs:
            text = normalize_text(paragraph.text)
            if not text:
                continue
            style_name = (paragraph.style.name or "").casefold()
            if style_name.startswith("heading"):
                suffix = style_name.removeprefix("heading").strip()
                level = int(suffix) if suffix.isdigit() else 1
                blocks.append(ParsedBlock(text, ChunkType.SECTION, heading_level=level))
                if level == 1 and title == path.stem:
                    title = text
            elif style_name.startswith("list"):
                blocks.append(ParsedBlock(text, ChunkType.LIST))
            else:
                blocks.append(ParsedBlock(text, ChunkType.PARAGRAPH))
        for table in document.tables:
            rows = [
                " | ".join(normalize_text(cell.text) for cell in row.cells) for row in table.rows
            ]
            text = normalize_text("\n".join(rows), preserve_lines=True)
            if text:
                blocks.append(ParsedBlock(text, ChunkType.TABLE))
        return ParsedDocument(title=title, blocks=blocks, metadata={"format": "docx"})
