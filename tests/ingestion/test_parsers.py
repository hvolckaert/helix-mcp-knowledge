from pathlib import Path

import pytest
from docx import Document as DocxDocument
from reportlab.pdfgen import canvas

from helix_mcp_knowledge.config import OcrSettings
from helix_mcp_knowledge.errors import IngestionError
from helix_mcp_knowledge.ingestion.chunker import Chunker
from helix_mcp_knowledge.ingestion.parsers.pdf import PdfParser
from helix_mcp_knowledge.ingestion.parsers.registry import ParserRegistry
from helix_mcp_knowledge.models.chunk import ChunkType
from helix_mcp_knowledge.ocr_component import OcrPage


@pytest.fixture
def parsers() -> ParserRegistry:
    return ParserRegistry(
        allowed_extensions=[".pdf", ".docx", ".md", ".html", ".txt"],
        max_file_size_mb=5,
    )


def test_markdown_preserves_headings_lists_tables_and_code(
    tmp_path: Path, parsers: ParserRegistry
) -> None:
    path = tmp_path / "design.md"
    path.write_text(
        """# example_project Design

## Reconciliation

The precedence rule uses BMC_ComputerSystem.

- Dataset A
- Dataset B

| Class | Value |
| --- | --- |
| ComputerSystem | Primary |

```sql
SELECT * FROM AST:Attributes;
```
""",
        encoding="utf-8",
    )
    parsed = parsers.parse(path)
    assert parsed.title == "example_project Design"
    assert {block.chunk_type for block in parsed.blocks} >= {
        ChunkType.SECTION,
        ChunkType.PARAGRAPH,
        ChunkType.LIST,
        ChunkType.TABLE,
        ChunkType.CODE,
    }


def test_html_strips_active_content_and_preserves_structure(
    tmp_path: Path, parsers: ParserRegistry
) -> None:
    path = tmp_path / "guide.html"
    path.write_text(
        """
        <html><head><title>CMDB Guide</title><script>secret()</script></head>
        <body><h1>Classes</h1><p>BMC_ComputerSystem details.</p>
        <ul><li>Item one</li><li>Item two</li></ul></body></html>
        """,
        encoding="utf-8",
    )
    parsed = parsers.parse(path)
    text = "\n".join(block.text for block in parsed.blocks)
    assert parsed.title == "CMDB Guide"
    assert "secret" not in text
    assert "BMC_ComputerSystem" in text
    assert any(block.chunk_type is ChunkType.LIST for block in parsed.blocks)


def test_docx_extracts_headings_paragraphs_and_tables(
    tmp_path: Path, parsers: ParserRegistry
) -> None:
    path = tmp_path / "procedure.docx"
    document = DocxDocument()
    document.core_properties.title = "Reconciliation Procedure"
    document.add_heading("Preparation", level=1)
    document.add_paragraph("Validate the dataset before execution.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Dataset"
    table.cell(0, 1).text = "Priority"
    document.save(path)

    parsed = parsers.parse(path)
    assert parsed.title == "Reconciliation Procedure"
    assert any(block.heading_level == 1 for block in parsed.blocks)
    assert any(block.chunk_type is ChunkType.TABLE for block in parsed.blocks)


def test_docx_replaces_generic_core_title_with_first_heading(
    tmp_path: Path, parsers: ParserRegistry
) -> None:
    path = tmp_path / "example_component.docx"
    document = DocxDocument()
    document.core_properties.title = "Word Document"
    document.add_heading("example_component Integration", level=1)
    document.add_paragraph("Integration details.")
    document.save(path)

    parsed = parsers.parse(path)

    assert parsed.title == "example_component Integration"


def test_pdf_extracts_each_text_page(tmp_path: Path, parsers: ParserRegistry) -> None:
    path = tmp_path / "official.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.setTitle("Official CMDB")
    pdf.drawString(72, 760, "BMC Helix CMDB reconciliation guidance")
    pdf.showPage()
    pdf.drawString(72, 760, "ARERR 120029 troubleshooting")
    pdf.save()

    parsed = parsers.parse(path)
    assert parsed.title == "Official CMDB"
    assert parsed.metadata["pages"] == 2
    text = "\n".join(block.text for block in parsed.blocks)
    assert "BMC Helix CMDB" in text
    assert "ARERR 120029" in text


def test_image_only_pdf_is_rejected_without_ocr(tmp_path: Path, parsers: ParserRegistry) -> None:
    path = tmp_path / "blank.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.showPage()
    pdf.save()
    with pytest.raises(IngestionError, match="OCR"):
        parsers.parse(path)


def test_pdf_uses_ocr_only_for_pages_below_the_text_threshold(tmp_path: Path) -> None:
    path = tmp_path / "mixed.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(72, 760, "This native text page is long enough to avoid OCR processing.")
    pdf.showPage()
    pdf.showPage()
    pdf.save()

    class FakeOcr:
        def __init__(self) -> None:
            self.calls: list[list[int]] = []

        def extract(self, source: Path, page_numbers: list[int]):
            assert source == path
            self.calls.append(page_numbers)
            return {2: OcrPage(2, "Recognized scanned DEVICE design", 0.94)}

    ocr = FakeOcr()
    parser = PdfParser(
        ocr_settings=OcrSettings(enabled=True, min_text_characters=40),
        ocr=ocr,
    )

    parsed = parser.parse(path, allow_ocr=True)

    assert ocr.calls == [[2]]
    assert "Recognized scanned DEVICE design" in "\n".join(block.text for block in parsed.blocks)
    assert parsed.metadata["ocr"] == {
        "used": True,
        "engine": "rapidocr",
        "pages": [2],
        "average_confidence": 0.94,
    }


def test_official_pdf_scope_never_invokes_optional_ocr(tmp_path: Path) -> None:
    path = tmp_path / "official-scan.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.showPage()
    pdf.save()

    class UnexpectedOcr:
        def extract(self, source: Path, page_numbers: list[int]):
            raise AssertionError("official sources must not invoke project OCR")

    parser = PdfParser(
        ocr_settings=OcrSettings(enabled=True),
        ocr=UnexpectedOcr(),
    )

    with pytest.raises(IngestionError, match="optional project PDF OCR"):
        parser.parse(path, allow_ocr=False)


def test_chunker_preserves_heading_path_and_token_limit(
    tmp_path: Path, parsers: ParserRegistry
) -> None:
    path = tmp_path / "long.md"
    path.write_text(
        "# Design\n\n## Mapping\n\n" + "technical mapping value. " * 100,
        encoding="utf-8",
    )
    chunks = Chunker(target_tokens=30, max_tokens=40, overlap_tokens=5).chunk(parsers.parse(path))
    assert len(chunks) > 1
    assert all(chunk.heading_path == ["Design", "Mapping"] for chunk in chunks)
    assert all(chunk.token_count <= 40 for chunk in chunks)


def test_symbolic_link_source_is_rejected(tmp_path: Path, parsers: ParserRegistry) -> None:
    target = tmp_path / "target.txt"
    target.write_text("content", encoding="utf-8")
    link = tmp_path / "link.txt"
    try:
        link.symlink_to(target)
    except OSError as exc:
        pytest.skip(f"symbolic links are unavailable for this Windows runner: {exc}")
    with pytest.raises(IngestionError, match="symbolic-link"):
        parsers.parse(link)
